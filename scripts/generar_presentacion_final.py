"""Genera la Presentación Final del Producto (versión venta) en PPTX + PDF.

6 slides, formato HELIOS, dividido por integrante. Basada en el guion de venta
+ mejoras aprobadas por el PO (8-jul-2026):
  - El límite de citas por médico pasa de "roadmap" a "en producción" (F-01).
  - Número alineado al del médico real: 15 citas/día.
  - Métricas al día: 141 tests + CI/CD (auto-deploy con GitHub Actions).
  - La demo incluye el "Día completo" (el límite en acción).

Reparto: Marian -> Ethan -> José -> Alejandro -> Ashley -> Marian.

Output:
  dailys/Entrega_Viernes/Presentacion_Final/HELIOS_Presentacion_Final.pptx
  dailys/Entrega_Viernes/Presentacion_Final/HELIOS_Presentacion_Final.pdf
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "dailys" / "Entrega_Viernes" / "Presentacion_Final"
OUT_PPTX = OUT_DIR / "HELIOS_Presentacion_Final.pptx"
OUT_PDF = OUT_DIR / "HELIOS_Presentacion_Final.pdf"
QR_IMG = ROOT / "dailys" / "qr_helios_demo.png"
LOGO = ROOT / "frontend" / "public" / "helios-logo.png"
OUT_DIR.mkdir(parents=True, exist_ok=True)

# Paleta HELIOS
TEAL = RGBColor(0x0A, 0x84, 0x78)
TEAL_DARK = RGBColor(0x0A, 0x4F, 0x6B)
TEAL_LIGHT = RGBColor(0x1F, 0xA8, 0x9B)
ACCENT = RGBColor(0x3D, 0xDC, 0x97)
INK = RGBColor(0x0F, 0x2A, 0x36)
INK_MUTED = RGBColor(0x4A, 0x6B, 0x7A)
MIST = RGBColor(0xF4, 0xF7, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xF3, 0x4B, 0x56)      # rojo de marca HELIOS (logo)
GREEN_OK = RGBColor(0x10, 0xB9, 0x81)

TOTAL = 6


def tb(slide, l, t, w, h, text, *, size=14, bold=False, color=INK,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return box


def rect(slide, l, t, w, h, color, *, line=None, line_w=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def chip(slide, l, t, text, color=TEAL, w=1.8):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(w), Inches(0.38))
    c.fill.solid()
    c.fill.fore_color.rgb = color
    c.line.fill.background()
    c.shadow.inherit = False
    p = c.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return c


def foot(slide, text, n):
    tb(slide, Inches(0.5), Inches(7.05), Inches(10), Inches(0.3), text, size=9, color=INK_MUTED)
    tb(slide, Inches(12.0), Inches(7.05), Inches(1), Inches(0.3), f"{n} / {TOTAL}",
       size=9, color=INK_MUTED, align=PP_ALIGN.RIGHT)


def bullets(slide, l, t, w, h, items, *, size=14, color=INK, gap=8):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + it
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = "Calibri"
        p.space_after = Pt(gap)


def title(slide, chip_text, h1, *, chip_color=TEAL, chip_w=1.9, subtitle=None):
    rect(slide, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    chip(slide, Inches(0.5), Inches(0.4), chip_text, color=chip_color, w=chip_w)
    tb(slide, Inches(0.5), Inches(0.95), Inches(12.4), Inches(0.9), h1, size=28, bold=True, color=INK)
    if subtitle:
        tb(slide, Inches(0.5), Inches(1.62), Inches(12.4), Inches(0.4), subtitle, size=14, color=INK_MUTED)


# === SLIDE 1: Portada + gancho (Marian, 20s) ===
def s1():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, prs.slide_width, prs.slide_height, MIST)
    rect(slide, 0, 0, Inches(0.4), prs.slide_height, RED)
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(1.0), Inches(0.8), height=Inches(0.9))
    tb(slide, Inches(2.05), Inches(0.85), Inches(9), Inches(0.6), "HELIOS", size=40, bold=True, color=RED)
    tb(slide, Inches(2.07), Inches(1.55), Inches(9), Inches(0.4), "Citas médicas unificadas", size=15, color=INK_MUTED)

    tb(slide, Inches(1.0), Inches(2.55), Inches(11.6), Inches(0.8), "Presentación Final del Producto", size=32, bold=True, color=INK)

    rect(slide, Inches(1.0), Inches(3.5), Inches(11.4), Inches(1.7), WHITE, line=RED, line_w=1.5)
    tb(slide, Inches(1.25), Inches(3.68), Inches(11.0), Inches(1.4),
       "En Panamá, agendar una cita médica todavía significa llamar, esperar o hacer fila. "
       "HELIOS lo resuelve: agendar, cancelar, reprogramar y recibir aviso — en menos de 2 minutos, desde el celular.",
       size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE)

    tb(slide, Inches(1.0), Inches(5.4), Inches(11.6), Inches(0.5),
       "Hoy no solo les contamos cómo lo construimos: se los mostramos funcionando, y qué dijo un médico real que lo probó.",
       size=14, italic=True, color=TEAL_DARK)

    tb(slide, Inches(1.0), Inches(6.25), Inches(11.6), Inches(0.4), "Grupo HELIOS", size=13, bold=True, color=INK)
    tb(slide, Inches(1.0), Inches(6.6), Inches(11.8), Inches(0.4), "Jose Brown · Ethan Gonzalez · Alejandro Saenz · Ashley Martinez · Marian Barba", size=11.5, color=INK_MUTED)
    tb(slide, Inches(1.0), Inches(6.95), Inches(11), Inches(0.4), "UTP · Calidad de Software · 2026", size=10.5, color=INK_MUTED)
    tb(slide, Inches(12.0), Inches(7.05), Inches(1), Inches(0.3), f"1 / {TOTAL}", size=9, color=INK_MUTED, align=PP_ALIGN.RIGHT)


# === SLIDE 2: Demo en vivo (Ethan, 60s) ===
def s2():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "DEMO EN VIVO", "Primero, véanlo funcionar", chip_color=RED, chip_w=1.7,
          subtitle="Sin truco: no recargamos la página ni revisamos nada a mano")
    if QR_IMG.exists():
        slide.shapes.add_picture(str(QR_IMG), Inches(0.8), Inches(2.2), width=Inches(3.7), height=Inches(3.7))
        tb(slide, Inches(0.8), Inches(6.0), Inches(3.7), Inches(0.3), "Escanéalo y pruébalo", size=11, color=INK_MUTED, align=PP_ALIGN.CENTER)

    rect(slide, Inches(4.9), Inches(2.2), Inches(7.9), Inches(0.8), TEAL_DARK)
    tb(slide, Inches(4.9), Inches(2.3), Inches(7.9), Inches(0.6), "helios-web-production.up.railway.app",
       size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    pasos = [
        ("1", "Registrar paciente con su correo"),
        ("2", "Agendar una cita en segundos"),
        ("3", "Llega solo: campanita en la web + correo"),
        ("4", "Bonus: al llenar un médico, el día sale 'Día completo'"),
    ]
    for i, (n, p) in enumerate(pasos):
        y = Inches(3.35) + Inches(0.82) * i
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.9), y, Inches(0.5), Inches(0.5))
        circ.fill.solid(); circ.fill.fore_color.rgb = TEAL; circ.line.fill.background(); circ.shadow.inherit = False
        pp = circ.text_frame.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
        rr = pp.add_run(); rr.text = n; rr.font.size = Pt(15); rr.font.bold = True; rr.font.color.rgb = WHITE
        tb(slide, Inches(5.6), y + Inches(0.03), Inches(7.2), Inches(0.6), p, size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    foot(slide, "Ethan (Frontend) — 60s — Demo en vivo", 2)


# === SLIDE 3: Validación real (José, 60s) ===
def s3():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "VALIDACIÓN REAL", "Un médico real lo probó", chip_color=TEAL, chip_w=2.0,
          subtitle="Lo que valida un producto de salud no es el equipo — es alguien externo con criterio clínico")
    bullets(slide, Inches(0.6), Inches(2.35), Inches(12.2), Inches(2.2), [
        "Lo llevamos a un médico de salud pública (policlínica) — el contexto real para el que lo diseñamos",
        "Prueba guiada, sin pacientes ni consultas reales de por medio",
        "Su reacción: le pareció increíble, y destacó el diseño visual y los colores",
        "En salud eso importa: un sistema confuso genera desconfianza antes de usar una sola función",
    ], size=15, gap=11)

    rect(slide, Inches(0.5), Inches(5.15), Inches(12.3), Inches(1.5), MIST, line=TEAL, line_w=1.5)
    tb(slide, Inches(0.75), Inches(5.3), Inches(11.9), Inches(0.4), "Lo más valioso no fue el elogio — fue su recomendación:", size=14, bold=True, color=INK)
    tb(slide, Inches(0.75), Inches(5.72), Inches(11.9), Inches(0.8),
       "“Establezcan un máximo de 15 pacientes por día por médico.”  Esa cifra no salió de nosotros: "
       "salió de alguien que gestiona consultas reales todos los días.",
       size=14, italic=True, color=TEAL_DARK, anchor=MSO_ANCHOR.MIDDLE)
    foot(slide, "Jose (PO) — 60s — Validación real", 3)


# === SLIDE 4: De la recomendación a producción (Alejandro, 45s) ===
def s4():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "DE LA RECOMENDACIÓN → PRODUCCIÓN", "No lo dejamos en el roadmap", chip_color=RED, chip_w=3.2,
          subtitle="La recomendación del médico se volvió una función real en el mismo ciclo")

    rect(slide, Inches(0.5), Inches(2.25), Inches(12.3), Inches(1.15), TEAL_DARK)
    tb(slide, Inches(0.7), Inches(2.38), Inches(11.9), Inches(0.9),
       "El límite de 15 citas/día por médico que recomendó → lo implementamos, lo verificamos y YA está en producción.",
       size=16, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

    tb(slide, Inches(0.6), Inches(3.65), Inches(12.2), Inches(0.4), "Cómo lo blindamos (3 capas):", size=13, bold=True, color=INK)
    bullets(slide, Inches(0.7), Inches(4.1), Inches(12.1), Inches(1.5), [
        "Base de datos: rechaza la cita que exceda el tope, con candado de concurrencia (dos pacientes no lo saltan)",
        "Reprogramar respeta el tope del día destino",
        "La interfaz muestra los días llenos como 'Día completo'",
        "Verificado con una revisión adversarial de 5 lentes antes de desplegar",
    ], size=13.5, gap=6)

    rect(slide, Inches(0.5), Inches(6.0), Inches(12.3), Inches(0.75), MIST, line=GREEN_OK)
    tb(slide, Inches(0.7), Inches(6.1), Inches(11.9), Inches(0.6),
       "Y hoy, no en el roadmap: las notificaciones (campanita + correo) corren solas — nadie las reinicia ni las tiene en su laptop.",
       size=12.5, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    foot(slide, "Alejandro (Tech Lead) — 45s", 4)


# === SLIDE 5: Calidad + ingeniería (Ashley, 30s) ===
def s5():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "CALIDAD E INGENIERÍA", "No es solo una demo bien montada", chip_w=2.6,
          subtitle="La misma disciplina con la que lo estamos presentando")
    mets = [
        ("141", "pruebas automáticas en cada cambio"),
        (">60%", "cobertura mínima (rechaza el código si baja)"),
        ("2", "robots de auto-deploy (CI/CD, GitHub Actions)"),
        ("RLS", "cada paciente y médico solo ve LO SUYO"),
    ]
    for i, (num, lbl) in enumerate(mets):
        col = i % 2
        row = i // 2
        left = Inches(0.5) + Inches(6.25) * col
        yy = Inches(2.35) + Inches(1.5) * row
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, yy, Inches(6.0), Inches(1.3))
        card.fill.solid(); card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = TEAL; card.line.width = Pt(1.4); card.shadow.inherit = False
        tb(slide, left + Inches(0.25), yy + Inches(0.1), Inches(2.3), Inches(1.1), num, size=34, bold=True, color=TEAL, anchor=MSO_ANCHOR.MIDDLE)
        tb(slide, left + Inches(2.5), yy + Inches(0.1), Inches(3.3), Inches(1.1), lbl, size=13, color=INK, anchor=MSO_ANCHOR.MIDDLE)

    rect(slide, Inches(0.5), Inches(5.6), Inches(12.3), Inches(0.9), MIST, line=TEAL)
    tb(slide, Inches(0.7), Inches(5.72), Inches(11.9), Inches(0.7),
       "Herramientas: Vitest · SonarCloud · Supabase (Auth + RLS) · GitHub Actions · Railway.",
       size=13, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    foot(slide, "Ashley (QA) — 30s", 5)


# === SLIDE 6: Cierre (Marian, 30s) ===
def s6():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "CIERRE", "Ya no es un ejercicio de clase", chip_w=1.4)
    bullets(slide, Inches(0.7), Inches(2.25), Inches(12.0), Inches(2.2), [
        "Un sistema DESPLEGADO, en producción",
        "PROBADO por alguien ajeno al equipo (un médico real)",
        "Funcionando de forma AUTÓNOMA (notificaciones que corren solas)",
        "Resuelve un problema real en Panamá: agendar una cita no debería tomar más de 2 minutos",
    ], size=16, gap=13)

    rect(slide, Inches(0.5), Inches(5.15), Inches(12.3), Inches(1.35), RED)
    tb(slide, Inches(0.5), Inches(5.3), Inches(12.3), Inches(0.5), "Pruébenlo ustedes mismos ahora:", size=14, color=WHITE, align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.5), Inches(5.72), Inches(12.3), Inches(0.5), "helios-web-production.up.railway.app", size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb(slide, Inches(0.5), Inches(6.62), Inches(12.3), Inches(0.4), "Gracias.  ¿Preguntas?", size=15, bold=True, color=INK, align=PP_ALIGN.CENTER)
    foot(slide, "Marian (Scrum Master) — 30s", 6)


def to_pdf(pptx, pdf):
    try:
        import comtypes.client
    except ImportError:
        print("[warn] comtypes no instalado -> solo PPTX")
        return False
    app = None
    try:
        app = comtypes.client.CreateObject("PowerPoint.Application")
        app.Visible = 1
        deck = app.Presentations.Open(str(pptx.absolute()), WithWindow=False)
        deck.SaveAs(str(pdf.absolute()), 32)
        deck.Close()
        return True
    except Exception as e:
        print(f"[error] PowerPoint: {e}")
        return False
    finally:
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def main():
    s1(); s2(); s3(); s4(); s5(); s6()
    prs.save(str(OUT_PPTX))
    print(f"[ok] PPTX: {OUT_PPTX} ({OUT_PPTX.stat().st_size/1024:.1f} KB)")
    if to_pdf(OUT_PPTX, OUT_PDF):
        print(f"[ok] PDF: {OUT_PDF} ({OUT_PDF.stat().st_size/1024:.1f} KB)")


if __name__ == "__main__":
    main()
